from io import BytesIO

from docx import Document
from openpyxl import Workbook

from app.memory.documents import extract_document


def test_extract_document_keeps_docx_source_metadata():
    document = Document()
    document.add_paragraph("Revenue grew by 12 percent.")
    stream = BytesIO()
    document.save(stream)

    sections = extract_document("annual-report.docx", stream.getvalue())

    assert sections[0].metadata["source"] == "annual-report.docx"
    assert sections[0].metadata["location"] == "document"
    assert "Revenue grew" in sections[0].text


def test_extract_document_keeps_xlsx_sheet_metadata():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Income statement"
    worksheet.append(["Year", "Revenue"])
    worksheet.append([2025, 120])
    stream = BytesIO()
    workbook.save(stream)

    sections = extract_document("financials.xlsx", stream.getvalue())

    assert sections[0].metadata["location"] == "sheet Income statement"
    assert "Year | Revenue" in sections[0].text
