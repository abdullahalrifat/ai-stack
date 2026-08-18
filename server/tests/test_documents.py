from io import BytesIO

from docx import Document
from openpyxl import Workbook

from app.memory.documents import extract_document, extraction_quality


def test_extract_document_keeps_docx_source_metadata():
    document = Document()
    document.add_paragraph("Revenue grew by 12 percent.")
    stream = BytesIO()
    document.save(stream)

    sections = extract_document("annual-report.docx", stream.getvalue())

    assert sections[0].metadata["source"] == "annual-report.docx"
    assert sections[0].metadata["location"] == "document"
    assert "Revenue grew" in sections[0].text


def test_extract_document_keeps_xlsx_sheet_metadata():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Income statement"
    worksheet.append(["Year", "Revenue"])
    worksheet.append([2025, 120])
    stream = BytesIO()
    workbook.save(stream)

    sections = extract_document("financials.xlsx", stream.getvalue())

    assert sections[0].metadata["location"] == "sheet Income statement"
    assert "Year | Revenue" in sections[0].text
    assert sections[0].metadata["table_row_count"] == 2
    assert sections[0].metadata["quality_score"] > 0.9


def test_extract_document_includes_docx_tables():
    document = Document()
    document.add_heading("ACTIVE ITEMS", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item A"
    table.rows[0].cells[1].text = "Open"
    stream = BytesIO()
    document.save(stream)

    sections = extract_document("items.docx", stream.getvalue())

    assert "Item A | Open" in sections[0].text
    assert sections[0].metadata["section_kind"] == "primary"


def test_extraction_quality_flags_empty_image_only_text():
    quality = extraction_quality("")

    assert quality["needs_ocr"] is True
    assert quality["extraction_status"] == "low_quality"
